from pathlib import Path
import csv, math, sys, os
os.environ['MPLCONFIGDIR']='/tmp/lab2-report-mpl'
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path('/home/evan/Desktop/NTUST/105_1/Class/Computer_Aided_Design_for_VLSI_System/Lab2')
OUT=ROOT/'M11502102_report/draft'; QA=OUT/'qa'
sys.path.insert(0,str(ROOT/'grader/grader'))
import verifier

def stages(group):
 return list(csv.DictReader(open(ROOT/f'M11502102_report/tables/{group}/stages.csv')))
main=stages('testcases'); hold=stages('holdout')

D=Document(); sec=D.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.45);sec.bottom_margin=Cm(1.4);sec.left_margin=Cm(1.6);sec.right_margin=Cm(1.6)
sec.footer_distance=Cm(.65)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=D.styles[name];s.font.name='Noto Sans CJK TC';s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK TC')
 s.paragraph_format.space_after=Pt(5)
N=D.styles['Normal'];N.font.size=Pt(10.5);N.paragraph_format.line_spacing=Pt(16)
N.paragraph_format.widow_control=True
D.styles['Title'].font.size=Pt(19);D.styles['Title'].font.bold=True
D.styles['Title'].paragraph_format.space_after=Pt(3)
D.styles['Subtitle'].font.size=Pt(9);D.styles['Subtitle'].font.italic=False
for h in ['Heading 1','Heading 2']:
 D.styles[h].font.size=Pt(12);D.styles[h].font.bold=True
 D.styles[h].paragraph_format.space_before=Pt(8);D.styles[h].paragraph_format.space_after=Pt(4)
 D.styles[h].paragraph_format.keep_with_next=True
D.styles['Caption'].font.size=Pt(8.5);D.styles['Caption'].font.italic=False;D.styles['Caption'].font.bold=False;D.styles['Caption'].paragraph_format.line_spacing=Pt(12)
D.core_properties.title='Lab2 時脈緩衝器插入研究報告'
D.core_properties.author='M11502102'
D.core_properties.subject='合法建構與模擬退火最佳化'

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('M11502102   ·   ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)

def p(txt,style=None,sz=None,after=None):
 q=D.add_paragraph(txt,style)
 if sz:
  for r in q.runs:r.font.size=Pt(sz)
 if after is not None:q.paragraph_format.space_after=Pt(after)
 return q

def heading(txt):return p(txt,'Heading 1')
def lead(label,txt):
 q=D.add_paragraph();q.add_run(label+' ').bold=True;q.add_run(txt);return q

def formula(txt):
 q=p(txt,sz=10);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
 q.paragraph_format.space_before=Pt(2);q.paragraph_format.space_after=Pt(6)

def table(headers,rows,widths,fs=9.0):
 t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=str(x)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if ri==0:
   repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for ci,cell in enumerate(row.cells):
   cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr()
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EDF3' if ri==0 else ('F6F8FA' if ri%2==0 else 'FFFFFF'));pr.append(shade)
   mar=OxmlElement('w:tcMar')
   for side,val in [('top','60'),('bottom','60'),('left','75'),('right','75')]:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   for q in cell.paragraphs:
    q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=Pt(12)
    for r in q.runs:r.font.size=Pt(fs);r.font.bold=(ri==0)
 return t

p('Lab2 時脈緩衝器插入研究報告','Title')
p('學號 M11502102　｜　Computer Aided Design for VLSI System　｜　2026 年 10 月 13 日','Subtitle')
p('本研究以合法時脈樹為起點，結合空間分群、模擬退火與增量更新，在 fanout 與總輸出線長限制下減少時脈偏移及緩衝器成本。五個版本於 27 筆測資共產生 135 份合法輸出；最終版相對基礎建構版的 Score 幾何平均比值，在既有與新增測資分別為 0.469 與 0.301。')
heading('一 問題定義與合法性')
p('以 Manhattan 距離 d(u,v) = |xᵤ − xᵥ| + |yᵤ − yᵥ| 表示邊長，sink 到達時間 Tᵢ 為 SRC 至該 sink 的路徑長度總和。本題不計 buffer 本身的延遲，也不需實際繞線；目標為最小化下式。[1]')
formula('Skew = T_max − T_min　　 Score = Σ buffer cost + w_skew × Skew')
p('每個父節點均須滿足 children 數 ≤ fanout 上限、所有輸出邊長總和 ≤ length 上限。座標不可重疊；buffer 採驗證器要求的嚴格晶片內部座標。每個 buffer／sink 恰有一個父節點且可由 SRC 到達，sink 為葉節點，buffer 不可閒置；輸出依 BFS 深度分層並連續編號。[1,2]')
heading('二 初始合法解建構')
lead('由下往上分群。','將所有 sink 視為尚未連接的子樹根；若 SRC 的 fanout 與總線長足以驅動全部根節點，便直接連接。否則以 Morton 編碼交錯座標位元排序，使空間鄰近節點較容易被分在同一群，再貪心擴群。')
lead('找位置與選型號。','先以各座標的中位數降低群內總線長，在附近搜尋未佔用的整數位置，並選擇同時滿足 fanout、length 的最低成本型號。最終版再以粗網格與逐步細搜，最小化「子樹 skew × w_skew + 型號成本」，以極小線長項打破同分。此為平衡啟發式，並不保證零 skew。')
lead('處理無法合併的情況。','若一輪沒有合併，便插入 relay buffer，使子樹根與 SRC 的距離縮短，再繼續分群。新父節點只連接既有的獨立子樹，因此不會形成環。成功建樹後完整重算合法性與 Score；位置搜尋有範圍及輪數上限，不能宣稱對所有可行輸入都必定找到解。')
heading('三 在合法前提下最佳化')
lead('局部操作與還原。','退火每次以 20% 插入繞路 buffer、30% 移動 buffer、30% 改接父節點、10% 平衡重放、10% 交換同祖父下兩個節點的父節點進行試探。改接後刪除空 buffer，並重新選擇最低成本合法型號；只檢查受影響父節點的負載、線長及座標，失敗即還原。改接範圍限制在祖父或叔伯 buffer，避免形成環。')
lead('退火與關鍵路徑。','使用 E = Score + 0.3 × w_skew × σ，其中 σ 為 sink 到達時間標準差；它能在最大／最小到達時間未變時提供搜尋方向。合法候選若 ΔE > 0，以 exp(−ΔE/T) 機率接受，溫度隨時間幾何遞減；30% 的節點選擇優先來自最早或最晚路徑。N ≤ 50 使用四段不同 seed 的搜尋，其餘一段。')
lead('增量評估。','各節點保存子樹最短／最長路徑 lo、hi、sink 數及到達時間的一、二階總和；修改只向祖先更新，摘要不變即停止。lo(u) = min[d(u,c)+lo(c)]，hi 同理取 max，故根節點可直接取得 skew。每次重算成本約為受影響祖先的 fanout 總和，避免掃描整棵樹；每段搜尋結束仍以完整評估核對，僅用合法且 Score 更低的結果替換保留解。[2]')

D.add_page_break()
heading('四 實驗設定與同測資階段比較')
p('既有 testcases 共 15 筆：001 為課程範例，其餘由測資產生器建立，N = 1 至 10,000；新增 holdout 共 12 筆，由 seed 2026 產生，涵蓋均勻／群聚分布、不同 SRC 位置、驅動限制與 w_skew。兩組皆屬本地實驗，新增組也曾參與開發比較，不視為獨立盲測。[3]')
p('採 C++17，最終版的最佳化時間預算為 min(40, 2 + N/250) 秒。以下使用保存的同一批實測紀錄，重新以課程 verifier 檢查五版共 135 份輸出，全部 VALID，且 Score、T_max、T_min 與紀錄一致。最終版單筆實測 2.01 至 40.04 秒，低於課程 60 秒限制；此時間含本機建樹與輸出，不保證其他機器也有相同耗時。[1-3]')
p('表 1　v5 同次執行的初始合法解 → 最佳化結果','Caption')
rows=[]
for r in main:
 rows.append([r['case'],f"{int(r['N']):,} / {r['w_skew']}",f"{int(r['S1 skew']):,} → {int(r['S2 skew']):,}",f"{int(r['S1 cost']):,} → {int(r['S2 cost']):,}",f"{int(r['S1 Score']):,} → {int(r['S2 Score']):,}",r['improve %']+'%'])
table(['測資','N / w_skew','Skew','Buffer 成本','Score','降幅'],rows,[1.05,2.25,3.15,3.9,4.05,1.35],fs=9)
p('降幅 = (初始 Score − 最終 Score) / 初始 Score。表中起點為 v5 的 stage1，與表 2 的 v1 基礎建構版不同；各階段均使用相同輸入。[3]',sz=8.5,after=6)
lead('高 skew 權重的收益。','010 的 skew 由 3,007 降至 232，下降 92.3%；成本由 18,414 降至 18,126，Score 由 48,484 降至 20,446，下降 57.8%。總改善 28,038 中，有 27,750 來自加權 skew 的降低，占 99.0%，說明本例主要受時序平衡驅動。')
lead('成本與偏移的取捨。','新增組 010（N = 2,000、w_skew = 10）的成本由 6,485 增至 7,115，但 skew 由 2,733 降至 237，Score 仍由 33,815 降至 9,485，下降 72.0%。因此適度增加繞路 buffer，可用較小成本換取更大的偏移改善。')
lead('整體階段效果。','以各筆「最終／初始 Score」取幾何平均，v5 最佳化後在既有組為 0.614、新增組為 0.510，即分別降低約 38.6% 與 49.0%。001 與單 sink 的 014 未改善；前者已達整數座標的 skew 下限 1，後者 skew 為 0，但這不等於已證明其成本最佳。')

D.add_page_break()
heading('五 版本貢獻與結果討論')
p('表 2 比較五版最終 Score。各組以 v1 為基準，計算 G = exp(mean(log(Score_v / Score_v1)))，越低越好；如此可避免大型測資的絕對分數主導平均。v5 的程式內容與目前提交資料夾 M11502102 一致。[2,3]')
p('表 2　各版本的 Score 幾何平均比值與搜尋預算','Caption')
table(['版本','主要方法','既有 15 筆','新增 12 筆','最佳化預算'],[
 ['v1','基礎合法建構','1.000','1.000','無'],
 ['v2','＋爬山式局部搜尋','0.630','0.452','min(20, 2+N/500) 秒'],
 ['v3','＋模擬退火與 σ','0.521','0.380','同 v2'],
 ['v4','＋增量更新與關鍵路徑','0.487','0.313','同 v2'],
 ['v5','＋平衡放置與交換父節點','0.469','0.301','min(40, 2+N/250) 秒'],
],[1.15,5.7,2.15,2.15,4.6],fs=8.8)
p('v5 相對 v1 的比值降低約 53.1%／69.9%；相對 v4 約降低 3.6%／3.7%。但 v5 同時改變建樹與時間預算，不能把差距全歸因於平衡／交換操作。逐筆仍有退步，例如既有 005 的 Score 由 v4 的 1,007 增至 v5 的 1,047；新增 004 由 150 增至 156。',after=4)
D.add_picture(str(QA/'arrival_comparison.png'),width=Cm(17.6))
D.paragraphs[-1].paragraph_format.space_after=Pt(1);D.paragraphs[-1].paragraph_format.line_spacing=1.0
p('圖 1　既有 004 的 50 個 sink 到達時間。v1 範圍為 91 至 722，v5 為 380 至 393；skew 由 631 降至 13，Score 由 1,756 降至 476。每點代表一個 sink，縱向錯開僅便於辨識；這是版本間比較，非表 1 的 v5 階段起點。[3]','Caption')
lead('增量更新提高搜尋量。','既有 011（10,000 sinks）在同為約 20 秒的紀錄中，v3 試探 1,708,032 次，v4 為 34,201,600 次，約 20.0 倍；Score 由 35,357 降至 33,248。v4 也加入關鍵路徑選點，因此這是整體版本效果，並非單獨隔離增量更新的消融實驗。')
lead('以鬆弛下界觀察改善空間。','設 N 為 sink 數、F_SRC 為 SRC fanout 上限，則 Σ_b(F_b−1) ≥ N−F_SRC，故成本下界為 ceil(max(0,N−F_SRC) × min_t[C_t/(F_t−1)])（取 F_t > 1）。另外路徑長度奇偶只由端點決定；若 sinks 的 x+y 奇偶不同，skew 至少為 1。成本下界加上 w_skew × skew 下界，即得 Score 下界，但忽略幾何與線長限制，未必可達。')
p('既有 011 的最終 Score／下界 = 32,808／27,496 ≈ 1.19，013 為 940／821 ≈ 1.14；相較之下，012 為 977／81 ≈ 12.06，顯示僅用 fanout 下界不足以描述長距離 relay 的必要成本。這些比值用於衡量鬆弛下界差距，不能當成已知最佳解的誤差。')
lead('限制與後續方向。','本結果為每版每筆一次保存紀錄，時間制退火受執行負載影響；後續應固定同一時間預算、重複多個 seed，並分別停用平衡與交換操作，量測平均與變異。現有程式週期性保存搜尋狀態，仍可能漏掉兩次保存之間的最佳候選，可改為依 Score 即時保存並評估額外成本。')
p('資料來源','Heading 2')
p('[1] Project02.pdf；Project2 Requirement.pdf。 [2] M11502102/src/main.cpp；grader/grader/verifier.py。 [3] M11502102_report/results/{testcases,holdout}、tables/{testcases,holdout} 與 scripts/；TUTORIAL.md。上述路徑皆相對於 Lab2。',sz=8,after=0)

# Remove inherited template rules and disable document grid snapping.
for style in D.styles:
 for child in list(style.element.iter(qn('w:pBdr'))):
  child.getparent().remove(child)
 if style.type == 1:
  pr=style.element.get_or_add_pPr()
  snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pr.append(snap)
for elem in list(D._element.iter(qn('w:docGrid'))):elem.getparent().remove(elem)
for elem in list(D._element.iter(qn('w:pBdr'))):elem.getparent().remove(elem)
D.save(OUT/'M11502102_初版報告.docx')
print(OUT/'M11502102_初版報告.docx')
